from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SOURCE = Path(r"C:\Users\Saman\Downloads\Rapport 2B.docx")
OUTPUT = Path(r"C:\Users\Saman\OneDrive\Skrivbord\Projektkurs-inhemsk-turism-github\output\Rapport 2B - slutversion.docx")


def replace_text(paragraph, old, new):
    if old not in paragraph.text:
        return False
    # These target paragraphs have uniform formatting. Preserve the first run's formatting.
    full = paragraph.text.replace(old, new)
    runs = paragraph.runs
    if runs:
        runs[0].text = full
        for run in runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(full)
    return True


def insert_paragraph_before(paragraph, text, style_source):
    new_p = OxmlElement("w:p")
    paragraph._p.addprevious(new_p)
    from docx.text.paragraph import Paragraph

    inserted = Paragraph(new_p, paragraph._parent)
    inserted.style = style_source.style
    inserted.paragraph_format.keep_with_next = True
    inserted.add_run(text)
    return inserted


doc = Document(SOURCE)

replacements = {
    "Fullständig rådata från användartesterna återfinns i bilaga C.":
        "Fullständig rådata från användartesterna återfinns i bilaga D.",
    "4.4 Fortsatt iteration - Sommar Uppdateringar":
        "4.4 Fortsatt iteration - sommaruppdateringar",
    "Denna utveckling genomfördes som en ny iteration/sprint":
        "Denna utveckling genomfördes som en ny iteration",
    "Figur 10. Gps-Funktionen": "Figur 10. GPS-funktionen",
    "Figur 14. Formulär för att skapa ny resa":
        "Figur 14. Formulär för att skapa en ny resa",
    "5.7 Mörkt & Ljust läge": "5.7 Mörkt och ljust läge",
    "Figur 24 - GPS meny": "Figur 24 - GPS-meny",
    "Bilaga C - Resultat användartester": "Bilaga D - Resultat användartester",
}

for paragraph in doc.paragraphs:
    for old, new in replacements.items():
        replace_text(paragraph, old, new)

# Add a separate appendix heading directly before the first summer-update figure.
gps_caption = next(p for p in doc.paragraphs if p.text.strip() == "Figur 24 - GPS-meny")
appendix_b = next(p for p in doc.paragraphs if p.text.strip() == "Bilaga B - High-fidelity prototyper")
insert_paragraph_before(gps_caption, "Bilaga C - Sommaruppdateringar", appendix_b)

# Ask Word-compatible applications to refresh fields (including the TOC) on open.
settings = doc.settings._element
update = settings.find(qn("w:updateFields"))
if update is None:
    update = OxmlElement("w:updateFields")
    settings.append(update)
update.set(qn("w:val"), "true")

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
