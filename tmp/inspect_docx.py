from docx import Document
from lxml import etree

path = r"C:\Users\Saman\Downloads\Rapport 2B.docx"
doc = Document(path)
print("paragraphs", len(doc.paragraphs), "sections", len(doc.sections), "inline_shapes", len(doc.inline_shapes))
for i, p in enumerate(doc.paragraphs):
    text = p.text.replace("\t", " <TAB> ").strip()
    xml = p._p.xml
    flags = []
    if "w:type=\"page\"" in xml:
        flags.append("PAGEBREAK")
    if "w:lastRenderedPageBreak" in xml:
        flags.append("RENDEREDBREAK")
    if "w:fldChar" in xml or "w:instrText" in xml:
        flags.append("FIELD")
    if "w:drawing" in xml or "w:pict" in xml:
        flags.append("IMAGE")
    if text or flags:
        print(f"{i:04d} | {p.style.name!r} | {','.join(flags)} | {text[:180]}")

print("\nBREAK DETAILS")
for i in list(range(275, 305)) + list(range(325, 395)):
    p = doc.paragraphs[i]
    breaks = p._p.xpath('.//w:br')
    pb = p._p.pPr is not None and p._p.pPr.pageBreakBefore is not None
    sect = p._p.pPr is not None and p._p.pPr.sectPr is not None
    if p.text.strip() or breaks or pb or sect or 'drawing' in p._p.xml:
        print(i, repr(p.text[:80]), p.style.name, 'breaks', [b.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}type') for b in breaks], 'pb', pb, 'sect', sect, 'drawing', 'drawing' in p._p.xml)
